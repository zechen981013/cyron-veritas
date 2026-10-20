"""托管估值表：生成（模拟器用）与解析（运营每天用）。

格式按托管行常见的「证券投资基金估值表」：表头若干行（产品名、估值日期、单位），然后是科目明细：
  科目代码 | 科目名称 | 数量 | 单位成本 | 成本 | 成本占净值% | 市价 | 市值 | 市值占净值% | 估值增值 | 停牌信息
科目：1002 银行存款、1021 结算备付金、1102 交易性金融资产（1102.01.01 上交所 A 股 / 1102.01.02 深交所 A 股，明细到证券代码）、
     1203 应收股利、2206 应付管理人报酬、2207 应付托管费、4001 实收基金；末尾是资产合计、负债合计、资产净值、单位净值、累计单位净值。
解析器只依赖「科目代码」表头和科目编码规则；不同托管行的列名差异在 COLUMN_ALIASES 里配置。
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER = ["科目代码", "科目名称", "数量", "单位成本", "成本", "成本占净值%", "市价", "市值", "市值占净值%", "估值增值", "停牌信息"]
COLUMN_ALIASES = {"科目代码": ["科目代码"], "科目名称": ["科目名称"], "数量": ["数量", "持仓数量"], "单位成本": ["单位成本"],
                  "成本": ["成本", "成本-本币"], "市价": ["市价", "行情", "行情收市价"], "市值": ["市值", "市值-本币"],
                  "市值占净值%": ["市值占净值%", "市值占净值(%)", "市值占比"], "估值增值": ["估值增值", "估值增值-本币"], "停牌信息": ["停牌信息", "停牌"]}
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
NUM, QTY, PRICE, PCT = "#,##0.00", "#,##0", "0.000", "0.00"


# ---------------- 生成 ----------------
def write(path: Path, product: dict, date, positions: list[dict], cash: float, settlement: float, receivable: list[dict],
          mgmt_payable: float, custody_payable: float, shares: float, acc_nav_unit: float | None = None) -> dict:
    """positions: [{code, name, exchange(SH/SZ), qty, cost, price, suspended}]；receivable: [{code, name, amount}]。返回汇总数。"""
    stock_mv = sum(p["qty"] * p["price"] for p in positions)
    recv = sum(r["amount"] for r in receivable)
    total_assets = cash + settlement + stock_mv + recv
    total_liab = mgmt_payable + custody_payable
    net = total_assets - total_liab
    unit = net / shares
    acc = acc_nav_unit if acc_nav_unit is not None else unit
    pct = lambda x: x / net * 100

    wb = Workbook()
    ws = wb.active
    ws.title = "估值表"
    ws["A1"] = f"{product['name']}_证券投资基金估值表"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"估值日期：{pd.Timestamp(date):%Y-%m-%d}"
    ws["F2"] = f"托管人：{product.get('custodian', '')}"
    ws["A3"] = "单位：元"
    r0 = 5
    for j, h in enumerate(HEADER, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0E7C66"), Alignment(horizontal="center")
    rows: list[list] = []
    rows.append(["1002", "银行存款", None, None, cash, pct(cash), None, cash, pct(cash), 0.0, None])
    rows.append(["1021", "结算备付金", None, None, settlement, pct(settlement), None, settlement, pct(settlement), 0.0, None])
    cost_all = sum(p["cost"] for p in positions)
    rows.append(["1102", "交易性金融资产", None, None, cost_all, pct(cost_all), None, stock_mv, pct(stock_mv), stock_mv - cost_all, None])
    rows.append(["1102.01", "股票投资", None, None, cost_all, pct(cost_all), None, stock_mv, pct(stock_mv), stock_mv - cost_all, None])
    for ex, sub, label in (("SH", "01", "上交所A股"), ("SZ", "02", "深交所A股")):
        ps = sorted([p for p in positions if p["exchange"] == ex], key=lambda p: p["code"])
        if not ps:
            continue
        c_, m_ = sum(p["cost"] for p in ps), sum(p["qty"] * p["price"] for p in ps)
        rows.append([f"1102.01.{sub}", label, None, None, c_, pct(c_), None, m_, pct(m_), m_ - c_, None])
        for p in ps:
            mv = p["qty"] * p["price"]
            rows.append([f"1102.01.{sub}.{p['code']}", p["name"], p["qty"], p["cost"] / p["qty"] if p["qty"] else 0, p["cost"],
                         pct(p["cost"]), p["price"], mv, pct(mv), mv - p["cost"], "停牌" if p.get("suspended") else None])
    if receivable:
        rows.append(["1203", "应收股利", None, None, recv, pct(recv), None, recv, pct(recv), 0.0, None])
        for r in receivable:
            rows.append([f"1203.01.{r['code']}", r["name"], None, None, r["amount"], pct(r["amount"]), None, r["amount"], pct(r["amount"]), 0.0, None])
    rows.append([None, "资产合计", None, None, None, None, None, total_assets, pct(total_assets), None, None])
    rows.append(["2206", "应付管理人报酬", None, None, mgmt_payable, pct(mgmt_payable), None, mgmt_payable, pct(mgmt_payable), None, None])
    rows.append(["2207", "应付托管费", None, None, custody_payable, pct(custody_payable), None, custody_payable, pct(custody_payable), None, None])
    rows.append([None, "负债合计", None, None, None, None, None, total_liab, pct(total_liab), None, None])
    rows.append(["4001", "实收基金", shares, None, shares, None, None, None, None, None, None])
    rows.append([None, "资产净值", None, None, None, None, None, net, 100.0, None, None])
    rows.append([None, "单位净值", None, None, None, None, round(unit, 4), None, None, None, None])
    rows.append([None, "累计单位净值", None, None, None, None, round(acc, 4), None, None, None, None])
    for i, row in enumerate(rows, start=r0 + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = THIN
            if j == 3:
                c.number_format = QTY
            elif j in (4, 7):
                c.number_format = PRICE
            elif j in (6, 9):
                c.number_format = PCT
            elif j in (5, 8, 10):
                c.number_format = NUM
            if row[0] and "." not in str(row[0]) or (row[0] is None and j == 2):
                c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJK", (22, 18, 14, 10, 18, 10, 10, 18, 10, 16, 8)):
        ws.column_dimensions[col].width = w
    ws.cell(row=len(rows) + r0 + 2, column=1, value="制表：模拟托管人估值核算部　复核：　　说明：本表为模拟数据，仅用于系统演示")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"total_assets": total_assets, "total_liab": total_liab, "net_assets": net, "unit_nav": round(unit, 4), "shares": shares,
            "stock_mv": stock_mv, "cash": cash + settlement, "receivable": recv, "fee_payable": total_liab}


# ---------------- 解析 ----------------
def _num(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace(",", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def parse(path: Path) -> dict:
    """读一张估值表 → {date, summary, holdings, receivable}。不依赖生成器，按表头和科目编码识别。"""
    ws = load_workbook(path, data_only=True).active
    vals = [[c.value for c in row] for row in ws.iter_rows()]
    date = None
    for row in vals[:6]:
        for v in row:
            m = re.search(r"(\d{4}-\d{2}-\d{2})", str(v or ""))
            if m and "估值日期" in str(v):
                date = pd.Timestamp(m.group(1)).date()
    hdr_i = next(i for i, row in enumerate(vals) if any(str(v).strip() == "科目代码" for v in row if v))
    hdr = [str(v).strip() if v else "" for v in vals[hdr_i]]
    col = {}
    for std, aliases in COLUMN_ALIASES.items():
        col[std] = next((hdr.index(a) for a in aliases if a in hdr), None)
    get = lambda row, k: row[col[k]] if col.get(k) is not None and col[k] < len(row) else None
    holds, recv, summary = [], [], {}
    for row in vals[hdr_i + 1:]:
        code, name = str(get(row, "科目代码") or "").strip(), str(get(row, "科目名称") or "").strip()
        m = re.fullmatch(r"1102\.01\.0([12])\.(\d{6})", code)
        if m:
            holds.append({"code": m.group(2), "name": name, "exchange": "SH" if m.group(1) == "1" else "SZ", "qty": _num(get(row, "数量")),
                          "unit_cost": _num(get(row, "单位成本")), "cost": _num(get(row, "成本")), "price": _num(get(row, "市价")),
                          "mv": _num(get(row, "市值")), "weight": (_num(get(row, "市值占净值%")) or 0) / 100,
                          "unrealized": _num(get(row, "估值增值")), "suspended": get(row, "停牌信息") or ""})
            continue
        m = re.fullmatch(r"1203\.01\.(\d{6})", code)
        if m:
            recv.append({"code": m.group(1), "name": name, "amount": _num(get(row, "市值"))})
            continue
        key = {"1002": "bank", "1021": "settlement", "1102": "stock_mv", "1203": "receivable", "2206": "mgmt_payable",
               "2207": "custody_payable", "4001": "shares"}.get(code) or {"资产合计": "total_assets", "负债合计": "total_liab",
                                                                         "资产净值": "net_assets", "单位净值": "unit_nav", "累计单位净值": "acc_nav"}.get(name)
        if key:
            v = _num(get(row, "数量")) if key == "shares" else (_num(get(row, "市价")) if key in ("unit_nav", "acc_nav") else _num(get(row, "市值")))
            summary[key] = v
    return {"date": date, "summary": summary, "holdings": holds, "receivable": recv, "file": str(path)}
